"""Excel Service for CODE MAZE using openpyxl.
Exports Product Reports and User Directories safely.
Per §6: NEVER export or store plaintext passwords or password hashes in Excel.
Only profile metadata (id, name, email/mobile, role, office, created_at, last_login).
"""

import io
from datetime import datetime
from typing import List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_reports_to_excel(reports_data: List[dict]) -> bytes:
    """Export compliance report list to styled Excel workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compliance Reports"

    # Navy Header Styling
    header_fill = PatternFill(start_color="12355B", end_color="12355B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    headers = [
        "Report No", "Date", "Product Name", "Manufacturer", "Category",
        "Verdict", "Score (%)", "Avg OCR Conf (%)", "Inspector", "Violations Count"
    ]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for r in reports_data:
        ws.append([
            r.get("report_number", ""),
            r.get("created_at", ""),
            r.get("product_name", "N/A"),
            r.get("manufacturer_name", "N/A"),
            r.get("category", "General"),
            r.get("verdict", "NEEDS_REVIEW"),
            r.get("compliance_score", 0.0),
            r.get("avg_ocr_confidence", 0.0),
            r.get("inspector_name", "Officer on Duty"),
            r.get("violations_count", 0),
        ])

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def export_users_to_excel(users_data: List[dict]) -> bytes:
    """Admin-only export of user profiles.
    Zero sensitive data (no password, hash, or secret tokens).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registered Users"

    header_fill = PatternFill(start_color="0E7490", end_color="0E7490", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    headers = [
        "User ID", "Name", "Email", "Mobile", "Role",
        "Designation", "Office", "District", "State", "Active", "Created At", "Last Login"
    ]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for u in users_data:
        ws.append([
            str(u.get("id", "")),
            u.get("name", ""),
            u.get("email", ""),
            u.get("mobile", ""),
            u.get("role", ""),
            u.get("designation", ""),
            u.get("office", ""),
            u.get("district", ""),
            u.get("state", ""),
            "Yes" if u.get("is_active") else "No",
            str(u.get("created_at", "")),
            str(u.get("last_login", "")),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
